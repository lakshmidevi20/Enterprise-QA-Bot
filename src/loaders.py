from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}

def _base_metadata(path: Path) -> dict:
    return {
        "source": path.name,
        "source_path": str(path.resolve()),
        "file_type": path.suffix.lower(),
    }

def load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    docs = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            metadata = _base_metadata(path)
            metadata["page"] = i
            docs.append(Document(page_content=text, metadata=metadata))
    return docs

def load_docx(path: Path) -> list[Document]:
    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    if not text.strip():
        return []
    return [Document(page_content=text, metadata=_base_metadata(path))]

def load_text(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    if not text.strip():
        return []
    return [Document(page_content=text, metadata=_base_metadata(path))]

def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix in {".txt", ".md"}:
        return load_text(path)
    return []

def load_documents(data_dir: str | Path) -> list[Document]:
    data_dir = Path(data_dir)
    if not data_dir.exists():
        raise FileNotFoundError(f"Data directory does not exist: {data_dir}")

    docs: list[Document] = []
    for path in sorted(data_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.extend(load_file(path))

    if not docs:
        raise RuntimeError(
            f"No supported documents found in {data_dir}. "
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return docs
